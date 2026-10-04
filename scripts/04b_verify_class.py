# 반 단위 KESS 원자료를 변환 경로와 별개로 다시 더해 docs/data/stats.json 합계와 대조한다(전남광주판)
import openpyxl, json, os, sys
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
stats = json.load(open(f'{ROOT}/docs/data/stats.json', encoding='utf-8'))
n = lambda c: str(c).replace('\n', '').replace(' ', '') if c is not None else ''
v = lambda x: int(float(x)) if x not in (None, '') and str(x).strip() not in ('', '-') else 0
bad = 0
for yi, y in enumerate(stats['years']):
    wb = openpyxl.load_workbook(f'{ROOT}/raw/kess_class_{y}.xlsx', read_only=True)
    ws = next(wb[s] for s in wb.sheetnames if '학생수' in s and '요약' not in s)
    it = ws.iter_rows(values_only=True)
    for row in it:
        if row and any(n(c) == '학교명' for c in row): h = [n(c) for c in row]; break
    I = h.index
    stu = cls = 0; per = {}
    for r in it:
        if not r or r[I('시도')] not in ('광주', '전남') or not r[I('학교명')]: continue
        g = ('광주 ' if r[I('시도')] == '광주' else '') + str(r[I('행정구역')]).strip()
        s, c = v(r[I('학생수_계')]), v(r[I('학급수')])
        stu += s; cls += c; per[g] = per.get(g, 0) + s
    mine = stats['total']['all'][yi]
    ok = (stu, cls) == (mine['students'], mine['classes'])
    for g, val in per.items():
        if stats['sigun'][g]['all'][yi]['students'] != val: ok = False; print('  시군 불일치', y, g, val, stats['sigun'][g]['all'][yi]['students'])
    print(y, '원자료 학생', stu, '학급', cls, '| 지도', mine['students'], mine['classes'], '| 시군구', len(per), '→', '일치' if ok else '불일치')
    bad += not ok
schools = json.load(open(f'{ROOT}/docs/data/schools.json', encoding='utf-8'))
s26 = sum(s['stu'][4] or 0 for s in schools)
print('학교 점 합계 2026', s26, '/ 통계', stats['total']['all'][4]['students'], '일치' if s26 == stats['total']['all'][4]['students'] else '불일치')
bad += s26 != stats['total']['all'][4]['students']
sys.exit(1 if bad else 0)
